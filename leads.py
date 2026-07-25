"""
Business Leads Scraper (OpenStreetMap)
--------------------------------------
Pulls business leads (name, address, phone, website, email, coordinates)
for any city + category using the free, legal OpenStreetMap Overpass API,
then cleans, dedupes and exports them to CSV + formatted Excel.

Usage:
    python leads.py --city "Lahore" --category restaurant
    python leads.py --city "London" --category dentist --limit 300
    python leads.py --city "Berlin" --tag "shop=bakery"        # any raw OSM tag

Categories built in: restaurant, cafe, dentist, pharmacy, hospital, hotel,
gym, school, salon, supermarket, bank, lawyer

Author: <your name> | github.com/<your-username>
"""

import argparse
import csv
import sys
import time
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OVERPASS_URL = "https://overpass-api.de/api/interpreter"
HEADERS = {"User-Agent": "PortfolioLeadsScraper/1.0 (+https://github.com/your-username)"}

# Friendly names -> OpenStreetMap tags
CATEGORY_TAGS = {
    "restaurant": 'amenity"="restaurant',
    "cafe": 'amenity"="cafe',
    "dentist": 'amenity"="dentist',
    "pharmacy": 'amenity"="pharmacy',
    "hospital": 'amenity"="hospital',
    "hotel": 'tourism"="hotel',
    "gym": 'leisure"="fitness_centre',
    "school": 'amenity"="school',
    "salon": 'shop"="hairdresser',
    "supermarket": 'shop"="supermarket',
    "bank": 'amenity"="bank',
    "lawyer": 'office"="lawyer',
}

FIELDS = ["name", "category", "phone", "website", "email", "address", "latitude", "longitude"]


def build_query(city: str, tag: str, limit: int) -> str:
    """Build an Overpass QL query: all nodes/ways with the tag inside the city area."""
    return f"""
    [out:json][timeout:90];
    area["name"="{city}"]->.searchArea;
    (
      node["{tag}"](area.searchArea);
      way["{tag}"](area.searchArea);
    );
    out center tags {limit};
    """


def fetch_leads(city: str, tag: str, category_label: str, limit: int) -> list[dict]:
    """Query the Overpass API and normalise the results into lead dicts."""
    query = build_query(city, tag, limit)
    print(f"Querying OpenStreetMap for '{category_label}' in {city} ...")
    resp = requests.post(OVERPASS_URL, data={"data": query}, headers=HEADERS, timeout=120)
    resp.raise_for_status()
    elements = resp.json().get("elements", [])

    leads = []
    for el in elements:
        tags = el.get("tags", {})
        name = tags.get("name", "").strip()
        if not name:
            continue  # unnamed places are not useful leads

        # Coordinates: nodes have lat/lon directly; ways have a 'center'
        lat = el.get("lat") or el.get("center", {}).get("lat", "")
        lon = el.get("lon") or el.get("center", {}).get("lon", "")

        address = ", ".join(
            part
            for part in [
                tags.get("addr:housenumber", ""),
                tags.get("addr:street", ""),
                tags.get("addr:city", ""),
                tags.get("addr:postcode", ""),
            ]
            if part
        )

        leads.append(
            {
                "name": name,
                "category": category_label,
                "phone": tags.get("phone") or tags.get("contact:phone", ""),
                "website": tags.get("website") or tags.get("contact:website", ""),
                "email": tags.get("email") or tags.get("contact:email", ""),
                "address": address,
                "latitude": lat,
                "longitude": lon,
            }
        )
    return leads


def clean_leads(leads: list[dict]) -> list[dict]:
    """Dedupe by (name, phone) and tidy up phone formatting."""
    seen = set()
    cleaned = []
    for lead in leads:
        key = (lead["name"].lower(), lead["phone"])
        if key in seen:
            continue
        seen.add(key)
        lead["phone"] = lead["phone"].replace(" ", " ").strip()
        cleaned.append(lead)
    return cleaned


def export_csv(leads: list[dict], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(leads)


def export_excel(leads: list[dict], path: Path, title: str) -> None:
    """Formatted, client-ready Excel sheet."""
    font_name = "Arial"
    header_fill = PatternFill("solid", fgColor="204E3D")
    header_font = Font(name=font_name, bold=True, color="FFFFFF", size=11)
    body_font = Font(name=font_name, size=10)
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"

    pretty_headers = ["Business Name", "Category", "Phone", "Website", "Email",
                      "Address", "Latitude", "Longitude"]
    ws.append(pretty_headers)
    for col in range(1, len(pretty_headers) + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for lead in leads:
        ws.append([lead[f] for f in FIELDS])

    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(FIELDS)):
        for cell in row:
            cell.font = body_font
            cell.border = border

    widths = {1: 34, 2: 14, 3: 18, 4: 34, 5: 28, 6: 38, 7: 11, 8: 11}
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:H{max(ws.max_row, 1)}"

    wb.save(path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Business leads scraper (OpenStreetMap)")
    parser.add_argument("--city", required=True, help='e.g. "Lahore", "London"')
    parser.add_argument("--category", choices=sorted(CATEGORY_TAGS), help="business type")
    parser.add_argument("--tag", help='raw OSM tag, e.g. "shop=bakery" (overrides --category)')
    parser.add_argument("--limit", type=int, default=500, help="max results")
    parser.add_argument("--out-dir", default="output", help="output folder")
    args = parser.parse_args()

    if args.tag:
        key, _, value = args.tag.partition("=")
        tag = f'{key}"="{value}'
        label = args.tag
    elif args.category:
        tag = CATEGORY_TAGS[args.category]
        label = args.category
    else:
        parser.error("provide --category or --tag")

    try:
        leads = fetch_leads(args.city, tag, label, args.limit)
    except requests.RequestException as exc:
        print(f"Request failed: {exc}")
        return 1

    leads = clean_leads(leads)
    if not leads:
        print("No named businesses found — try a bigger city or another category.")
        return 1

    out_dir = Path(args.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    slug = f"{args.city.lower().replace(' ', '_')}_{label.replace('=', '_')}"
    csv_path = out_dir / f"leads_{slug}.csv"
    xlsx_path = out_dir / f"leads_{slug}.xlsx"

    export_csv(leads, csv_path)
    export_excel(leads, xlsx_path, f"{label.title()} — {args.city}")

    with_phone = sum(1 for l in leads if l["phone"])
    with_web = sum(1 for l in leads if l["website"])
    print("\n===== SUMMARY =====")
    print(f"Total leads   : {len(leads)}")
    print(f"With phone    : {with_phone}")
    print(f"With website  : {with_web}")
    print(f"CSV           : {csv_path}")
    print(f"Excel         : {xlsx_path}")

    time.sleep(1)  # be gentle with the free API if scripted in a loop
    return 0


if __name__ == "__main__":
    sys.exit(main())
